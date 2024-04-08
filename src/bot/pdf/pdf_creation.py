import asyncio
import os
from docx import Document

from aiogram.types import FSInputFile

from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_RIGHT, TA_LEFT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from src.bot.kbds.text_builder import BUTTON_MENU
from src.bot.logic.user_private.fsm.flood_exp_branch.quiz.scenes_common import FloodExpAnswers
from src.configuration import conf

template_path = r'C:\08.04.2024\src\bot\common\form_template.docx'
font_path = r'C:\08.04.2024\src\bot\common\TimesNewRoman.ttf'

styles = getSampleStyleSheet() # дефолтовые стили
# the magic is here
styles['Normal'].fontName='TimesNewRoman'
styles['Heading1'].fontName='TimesNewRoman'

pdfmetrics.registerFont(TTFont('TimesNewRoman', font_path, 'UTF-8'))


style_right = ParagraphStyle(
        name='left',
        parent=styles['Normal'],
        fontSize=12,
        fontName='TimesNewRoman',
        spaceBefore=3,
        spaceAfter=3,
        alignment=TA_RIGHT)

style_left = ParagraphStyle(
        name='left',
        parent=styles['Normal'],
        fontSize=12,
        fontName='TimesNewRoman',
        alignment=TA_LEFT)

style_center = ParagraphStyle(
        name='left',
        parent=styles['Normal'],
        fontSize=12,
        fontName='TimesNewRoman',
        alignment=TA_CENTER)
def replace_placeholders(text, **kwargs):
    for key, value in kwargs.items():
        placeholder = "{" + key + "}"
        if placeholder in text:
            text = text.replace(placeholder, str(value))
    return text

async def send_pdf(user, flood_exp_answers, bot):
    target_path = rf'C:\08.04.2024\src\bot\pdf\{user.user_id}.pdf'
    doc = []
    docx = Document(template_path)
    zayava_flag = False
    fio = f'{flood_exp_answers.form[0]}'
    telephone = f'{flood_exp_answers.form[1]}'
    adress = f'Вулиця {flood_exp_answers.form[2]}, буд. №{flood_exp_answers.form[3]}, кв. №{flood_exp_answers.form[4]}'
    date = f'{flood_exp_answers.form[5]}'

    for line in docx.paragraphs:

        if zayava_flag:
            style = style_left
        elif 'Заява' == line.text:
            style = style_center
            zayava_flag = True
        else:
            style = style_right

        line.text =  replace_placeholders(line.text, fio=fio, telephone=telephone,
                                          adress=adress, date=date)

        if line.text == '':
            doc.append(Spacer(1, 12))
        doc.append(Paragraph(line.text, style))

    SimpleDocTemplate(filename=target_path,
                      rightMagrin=2.54 * cm, leftMagrin=2.54 * cm,
                      topMargin=2.54 * cm, bottomMagrin=2.54 * cm).build(doc)

    file = FSInputFile(target_path)

    message = await bot.send_document(chat_id=user.user_id, caption='Заявление принято, в ближайшее время специалист его'
                                                                ' проверит. А таĸже вы можете проверить статус модерации в меню.',
                                  document=file, reply_markup = BUTTON_MENU)

    #doc_id = doc.document.file_id

    os.remove(target_path)

    return message





